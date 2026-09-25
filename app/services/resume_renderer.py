from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.models.schemas import Resume


def render_resume_docx(resume: Resume) -> bytes:
    """
    Render a structured resume as a simple ATS-friendly DOCX document.
    """

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    _add_contact_header(document, resume)

    if resume.summary:
        _add_section_heading(document, "Summary")
        document.add_paragraph(resume.summary)

    if resume.skills:
        _add_section_heading(document, "Skills")
        document.add_paragraph(", ".join(resume.skills))

    if resume.experience:
        _add_section_heading(document, "Experience")
        for item in resume.experience:
            heading = _join_non_empty(
                [item.title, item.company],
                " | ",
            )
            dates = _join_non_empty(
                [item.start_date, item.end_date],
                " - ",
            )
            _add_entry_heading(document, heading, dates)
            _add_bullets(document, item.bullets)

    if resume.projects:
        _add_section_heading(document, "Projects")
        for project in resume.projects:
            if project.name:
                _add_entry_heading(document, project.name, "")
            if project.description:
                document.add_paragraph(project.description)
            _add_bullets(document, project.bullets)

    if resume.education:
        _add_section_heading(document, "Education")
        for item in resume.education:
            heading = _join_non_empty(
                [item.degree, item.school],
                " | ",
            )
            dates = _join_non_empty(
                [item.start_date, item.end_date],
                " - ",
            )
            _add_entry_heading(document, heading, dates)

    if resume.certifications:
        _add_section_heading(document, "Certifications")
        _add_bullets(document, resume.certifications)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_contact_header(document: Document, resume: Resume) -> None:
    contact = resume.contact
    name = contact.name or "Resume"
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(16)

    details = [
        contact.email,
        contact.phone,
        contact.location,
        contact.linkedin,
    ]
    detail_text = " | ".join(value for value in details if value)
    if detail_text:
        contact_paragraph = document.add_paragraph(detail_text)
        contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_section_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)


def _add_entry_heading(
    document: Document,
    heading: str,
    dates: str,
) -> None:
    if not heading and not dates:
        return

    paragraph = document.add_paragraph()
    run = paragraph.add_run(heading)
    run.bold = True
    if dates:
        paragraph.add_run(f"  {dates}")


def _add_bullets(
    document: Document,
    values: list[str],
) -> None:
    for value in values:
        if value.strip():
            document.add_paragraph(
                value,
                style="List Bullet",
            )


def _join_non_empty(
    values: list[str | None],
    separator: str,
) -> str:
    return separator.join(value for value in values if value)
